from pathlib import Path
from typing import List, Dict, Any

from openpyxl import load_workbook


class ExcelService:
    def load_workbook(self, file_path: Path):
        if file_path.suffix.lower() == ".csv":
            from openpyxl import Workbook
            workbook = Workbook()
            ws = workbook.active
            with open(file_path, "r", encoding="utf-8", newline="") as f:
                rows = [line.rstrip("\n").split(",") for line in f]
            for row in rows:
                ws.append(row)
            return workbook

        return load_workbook(file_path, data_only=True)

    def get_sheet_metadata(self, workbook):
        sheets = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            sample_rows = []
            for row in sheet.iter_rows(min_row=2, max_row=min(5, sheet.max_row), values_only=True):
                sample_rows.append(list(row))
            sheets.append({
                "name": sheet_name,
                "headers": headers,
                "rowCount": sheet.max_row,
                "sampleRows": sample_rows,
            })
        return sheets

    def read_rows(self, sheet):
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = list(rows[0])
        data = []
        for r in rows[1:]:
            record = dict(zip(headers, r))
            data.append(record)
        return data

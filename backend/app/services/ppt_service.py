from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any

from pptx import Presentation
from pptx.util import Inches, Pt


class PPTService:
    def generate_presentation(self, rows: List[Dict[str, Any]], selected_columns: List[Dict[str, str]], title: str, file_name: str):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        # Title slide
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Source file: {file_name}\nGenerated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"

        # Summary slide
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Data Summary"
        body = slide.shapes.placeholders[1].text_frame
        body.text = f"Rows processed: {len(rows)}\nColumns selected: {len(selected_columns)}"

        for idx, column in enumerate(selected_columns, start=1):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = column["name"]
            tf = slide.shapes.placeholders[1].text_frame
            tf.word_wrap = True
            tf.text = self._summarize_column(rows, column["name"], column["formula"])

        output_dir = Path(__file__).resolve().parent.parent.parent / "outputs"
        output_dir.mkdir(exist_ok=True)
        out_path = output_dir / f"{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{title.replace(' ', '_')}.pptx"
        prs.save(str(out_path))
        return out_path

    def _summarize_column(self, rows: List[Dict[str, Any]], column_name: str, formula: str):
        values = []
        for row in rows:
            value = row.get(column_name)
            if value is not None:
                try:
                    values.append(float(value))
                except (TypeError, ValueError):
                    continue

        if not values:
            return "No numeric data available"

        if formula == "subtraction":
            total = 0
            for i in range(1, len(values)):
                total += values[i] - values[i - 1]
            return f"Subtraction total: {total:.2f}"

        if formula == "multiplication":
            total = 1
            for value in values:
                total *= value
            return f"Multiplication result: {total:.2f}"

        if formula == "growth_percent":
            if len(values) < 2:
                return "Growth percentage requires at least 2 points"
            start = values[0]
            end = values[-1]
            pct = ((end - start) / start * 100) if start else 0
            return f"Growth: {pct:.2f}%"

        if formula == "percentage_of_total":
            total = sum(values)
            if not total:
                return "Total is zero"
            pct = (values[-1] / total) * 100 if values else 0
            return f"Share of total: {pct:.2f}%"

        return f"Average: {sum(values) / len(values):.2f}\nMin: {min(values):.2f}\nMax: {max(values):.2f}"
